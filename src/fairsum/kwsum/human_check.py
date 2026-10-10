"""Human check of the scorers on summary propositions (spec §5, human check protocol).

The sentiment and NLI scorers were validated on source reviews against star
ratings; summaries are a different register. This module draws a blinded
sample of summary propositions, writes an annotation sheet and a reading pack,
reads filled sheets back (refusing anything ambiguous), and compares human
labels with the scorer's.

Blinding: the sheet and the reading pack show the product, the 8 reviews and
the sentences to label. Model, condition, level and review ratings live only in
the key file, which is kept outside the annotation folder.
"""

from __future__ import annotations

import html
import json
import math
import random
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable, Mapping, Sequence

import numpy as np

from fairsum.annotation import krippendorff_alpha
from fairsum.kwsum.collections import Collection
from fairsum.kwsum.compliance import classify_output
from fairsum.kwsum.propositions import split_summary

CLAIM_VALUES = ("yes", "no")
SENTIMENT_VALUES = ("positive", "negative", "neutral", "mixed")
SUPPORT_VALUES = ("supported", "unsupported", "contradicted")
SHEET_COLUMNS = ("item_id", "sentence", "claim", "sentiment", "support", "notes")

GUIDE = [
    "HOW TO LABEL (about 90 minutes: 25 summaries x 4 sentences)",
    "",
    "1. Open reading_pack.html in a browser and your labels_<name>.xlsx in Excel, side by side.",
    "2. For each summary S01, S02, ...: read its 8 reviews, then label its 4 sentences in the sheet.",
    "3. Judge every sentence on its own. 'It' or 'this' means the product or drug named in the heading.",
    "4. Pick values from the drop-down lists. Save as you go; you can stop and continue later.",
    "",
    "claim",
    "  yes: the sentence says something about the product or about reviewers' experiences.",
    "  no: a heading, a note about the task ('Here is a summary', 'Let's count the reviews'),",
    "      a bare count ('6 reviews are positive'), or a fragment you cannot judge.",
    "      Leave sentiment and support blank when claim is no.",
    "",
    "sentiment: what the sentence says about the product",
    "  positive: praise, benefit, relief, satisfaction. 'It helped me sleep through the night.'",
    "  negative: complaint, side effect, failure, poor value. 'Several users reported weight gain.'",
    "  neutral: a fact with no judgement. 'It is taken once a day at bedtime.'",
    "  mixed: clearly both in one sentence. 'It worked well but caused nausea.'",
    "  Sentences about the overall opinion count too: 'received overwhelmingly positive",
    "      reviews' is positive; 'received mixed reviews' is mixed.",
    "",
    "support: check the sentence against all 8 reviews",
    "  supported: at least one review states it or clearly implies it.",
    "      Ignore quantity words ('many', 'most', 'some reviewers'): judge the content only.",
    "  contradicted: no review supports it and at least one review says the opposite.",
    "      Example: 'No side effects were reported' when a review describes side effects.",
    "  unsupported: no review says it or its opposite; the detail is not in the reviews.",
    "      Example: 'It is approved by the FDA' when no review mentions approval.",
    "  If one review supports the sentence and another disagrees, choose supported.",
    "  Exception: a sentence about the balance of opinion ('most reviews are positive',",
    "      'received mixed reviews') is judged against the 8 reviews as a whole:",
    "      supported if it fits, contradicted if it clearly does not.",
    "",
    "notes: optional. Write anything that made a label hard to choose.",
]


@dataclass(frozen=True)
class SampledSummary:
    summary_id: str
    collection_id: str
    model: str
    condition: str
    dataset: str
    level: str
    summary: str
    sentences: tuple[str, ...]


def item_ids(summaries: Sequence[SampledSummary]) -> list[str]:
    return [f"{s.summary_id}-{k}" for s in summaries for k in range(1, len(s.sentences) + 1)]


def sample_summaries(generations: Iterable[Mapping], n_summaries: int = 25, per_summary: int = 4,
                     seed: int = 42, max_new_tokens: int = 256) -> list[SampledSummary]:
    """Draw summaries evenly over (model, dataset), then `per_summary` propositions each.

    Only complete outputs (status `ok`) with enough propositions are eligible.
    Within a (model, dataset) stratum the draw prefers an unused collection,
    then the least-used condition, then the least-used level, so conditions
    and levels spread across the sample. Summary ids are assigned after a
    shuffle, so their order says nothing about model or condition.
    """
    eligible = []
    for g in generations:
        status = classify_output(g["raw"], int(g.get("completion_tokens", 0)), max_new_tokens,
                                 stop_reason=g.get("stop_reason"))
        if status != "ok":
            continue
        props = tuple(split_summary(g["raw"]).propositions)
        if len(props) >= per_summary:
            eligible.append((g, props))
    if len(eligible) < n_summaries:
        raise ValueError(f"only {len(eligible)} eligible summaries (status ok, at least "
                         f"{per_summary} propositions); {n_summaries} needed")

    # Sorted first, so the draw does not depend on the order rows were written in.
    eligible.sort(key=lambda e: (e[0]["model"], e[0]["dataset"], e[0]["condition"],
                                 e[0]["level"], e[0]["collection_id"]))
    rng = random.Random(seed)
    strata: dict[tuple[str, str], list] = defaultdict(list)
    for e in eligible:
        strata[(e[0]["model"], e[0]["dataset"])].append(e)
    keys = sorted(strata)
    rng.shuffle(keys)

    used_condition, used_level, used_collection = Counter(), Counter(), set()
    chosen = []
    while len(chosen) < n_summaries:
        for key in keys:
            if len(chosen) == n_summaries or not strata[key]:
                continue
            pool = [e for e in strata[key] if e[0]["collection_id"] not in used_collection] or strata[key]
            fewest = min(used_condition[e[0]["condition"]] for e in pool)
            pool = [e for e in pool if used_condition[e[0]["condition"]] == fewest]
            fewest = min(used_level[e[0]["level"]] for e in pool)
            pool = [e for e in pool if used_level[e[0]["level"]] == fewest]
            g, props = pick = rng.choice(pool)
            strata[key].remove(pick)
            used_condition[g["condition"]] += 1
            used_level[g["level"]] += 1
            used_collection.add(g["collection_id"])
            idx = sorted(rng.sample(range(len(props)), per_summary))
            chosen.append((g, tuple(props[i] for i in idx)))

    rng.shuffle(chosen)
    width = max(2, len(str(n_summaries)))
    return [SampledSummary(f"S{i:0{width}d}", g["collection_id"], g["model"], g["condition"],
                           g["dataset"], g["level"], g["raw"], sentences)
            for i, (g, sentences) in enumerate(chosen, start=1)]


def write_key(summaries: Sequence[SampledSummary], path: str | Path, seed: int) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps({
        "created_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "seed": seed,
        "summaries": [asdict(s) for s in summaries],
    }, indent=1, ensure_ascii=False), encoding="utf-8")


def read_key(path: str | Path) -> list[SampledSummary]:
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    return [SampledSummary(**{**s, "sentences": tuple(s["sentences"])}) for s in data["summaries"]]


def web_items(summaries: Sequence[SampledSummary], collections: Mapping[str, Collection]) -> list[dict]:
    """The study items for the web labelling page: product, 8 reviews and numbered sentences.

    As blind as the reading pack: no model, condition, level, collection id or review rating.
    """
    return [{"id": s.summary_id, "order": n, "topic": collections[s.collection_id].topic,
             "reviews": list(collections[s.collection_id].texts),
             "sentences": [{"item_id": f"{s.summary_id}-{k}", "text": text}
                           for k, text in enumerate(s.sentences, start=1)]}
            for n, s in enumerate(summaries, start=1)]


def web_labels(docs: Iterable[object]) -> dict[str, dict]:
    """Labels per item_id from the page's per-summary documents (`{"items": {item_id: {...}}}`).

    The documents are other people's input, so anything malformed is skipped and values
    are only normalised here; `read_sheet` does the validating once they are in a sheet.
    """
    labels: dict[str, dict] = {}
    for doc in docs:
        items = doc.get("items") if isinstance(doc, Mapping) else None
        if not isinstance(items, Mapping):
            continue
        for item, row in items.items():
            if not isinstance(row, Mapping):
                continue
            labels[str(item)] = {name: (_cell(row.get(name)).lower() or None)
                                 for name in ("claim", "sentiment", "support")}
            labels[str(item)]["notes"] = _cell(row.get("notes")) or None
    return labels


def write_sheet(summaries: Sequence[SampledSummary], path: str | Path,
                labels: Mapping[str, Mapping] | None = None) -> None:
    """A labelling sheet with drop-down lists, plus the guide on a second tab.

    Blank, or filled from `labels` (item_id -> claim/sentiment/support/notes), as the
    web page's labels are, so that `read_sheet` checks them like any hand-filled sheet."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "labels"
    ws.append(list(SHEET_COLUMNS))
    for s in summaries:
        for k, sentence in enumerate(s.sentences, start=1):
            item = f"{s.summary_id}-{k}"
            row = (labels or {}).get(item, {})
            ws.append([item, sentence, row.get("claim"), row.get("sentiment"), row.get("support"),
                       row.get("notes")])
    last = ws.max_row
    for name, values in (("claim", CLAIM_VALUES), ("sentiment", SENTIMENT_VALUES),
                         ("support", SUPPORT_VALUES)):
        letter = get_column_letter(SHEET_COLUMNS.index(name) + 1)
        dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True,
                            showErrorMessage=True, errorTitle="Not a label",
                            error="Pick one of: " + ", ".join(values))
        ws.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}{last}")
    widths = {"item_id": 9, "sentence": 90, "claim": 8, "sentiment": 11, "support": 14, "notes": 40}
    for i, name in enumerate(SHEET_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths[name]
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"

    guide = wb.create_sheet("guide")
    for line in GUIDE:
        guide.append([line])
    guide.column_dimensions["A"].width = 110

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


_PACK_CSS = """
body{font:16px/1.55 system-ui,-apple-system,Segoe UI,sans-serif;max-width:900px;margin:0 auto;
padding:24px 16px 80px;color:#1d1d1f;background:#fff}
h1{font-size:1.5rem}h2{font-size:1.2rem;margin-top:2.5rem;padding-top:1rem;border-top:2px solid #ddd}
h3{font-size:.95rem;text-transform:uppercase;letter-spacing:.04em;color:#555;margin-bottom:.4rem}
pre.guide{white-space:pre-wrap;background:#f5f5f7;padding:12px 14px;border-radius:8px;font-size:.9rem}
ol.items li{margin:.35rem 0;background:#fff8e1;padding:6px 10px;border-radius:6px;list-style:none}
ol.items b{font-family:ui-monospace,Consolas,monospace;margin-right:.5rem}
ol.reviews li{margin:.5rem 0}
nav{display:flex;flex-wrap:wrap;gap:.3rem .7rem}nav a{font-family:ui-monospace,Consolas,monospace}
@media (prefers-color-scheme:dark){body{background:#121212;color:#e8e8e8}
pre.guide{background:#1e1e1e}ol.items li{background:#3a3220}h2{border-color:#333}h3{color:#aaa}}
"""


def write_reading_pack(summaries: Sequence[SampledSummary], collections: Mapping[str, Collection],
                       path: str | Path) -> None:
    """One HTML page: the guide, then per summary its sentences to label and its 8 reviews.

    Review ratings are never shown, and neither is anything about how the
    summary was produced.
    """
    esc = html.escape
    parts = ["<!doctype html><html lang='en'><head><meta charset='utf-8'>",
             "<meta name='viewport' content='width=device-width,initial-scale=1'>",
             f"<title>Human check reading pack</title><style>{_PACK_CSS}</style></head><body>",
             "<h1>Human check: reading pack</h1>",
             f"<pre class='guide'>{esc(chr(10).join(GUIDE))}</pre>",
             "<nav>" + "".join(f"<a href='#{s.summary_id}'>{s.summary_id}</a>" for s in summaries)
             + "</nav>"]
    for s in summaries:
        c = collections[s.collection_id]
        parts.append(f"<section id='{s.summary_id}'><h2>{s.summary_id} &middot; {esc(c.topic)}</h2>")
        parts.append("<h3>Sentences to label</h3><ol class='items'>")
        parts += [f"<li><b>{s.summary_id}-{k}</b>{esc(sent)}</li>"
                  for k, sent in enumerate(s.sentences, start=1)]
        parts.append("</ol><h3>The 8 reviews</h3><ol class='reviews'>")
        parts += [f"<li>{esc(text)}</li>" for text in c.texts]
        parts.append("</ol></section>")
    parts.append("</body></html>")
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(parts), encoding="utf-8")


class HumanCheckSheetError(ValueError):
    """Every problem found in a filled sheet, so it can be fixed in one pass."""

    def __init__(self, problems: Sequence[str]):
        self.problems = list(problems)
        super().__init__(f"{len(self.problems)} problem(s) in the sheet:\n  "
                         + "\n  ".join(self.problems))


def _cell(value: object) -> str:
    return "" if value is None else str(value).strip()


def read_sheet(path: str | Path, expected_ids: Iterable[str] | None = None) -> dict[str, dict]:
    """Labels per item_id; claim = no leaves sentiment and support as None.

    A blank row is "not labelled", never a default label. Values are matched
    case- and space-insensitively. Problems are collected and raised together.
    """
    from openpyxl import load_workbook

    ws = load_workbook(path, read_only=True, data_only=True)["labels"]
    rows = list(ws.iter_rows(values_only=True))
    header = [_cell(h).lower() for h in rows[0]] if rows else []
    missing = [c for c in ("item_id", "claim", "sentiment", "support") if c not in header]
    if missing:
        raise HumanCheckSheetError([f"missing column(s): {', '.join(missing)}"])
    col = {name: header.index(name) for name in ("item_id", "claim", "sentiment", "support")}

    labels: dict[str, dict] = {}
    problems: list[str] = []
    seen: set[str] = set()
    for r, row in enumerate(rows[1:], start=2):
        item = _cell(row[col["item_id"]])
        if not item:
            if any(_cell(v) for v in row):
                problems.append(f"row {r}: missing item_id")
            continue
        where = f"{item} (row {r})"
        if item in seen:
            problems.append(f"{where}: duplicate item_id")
            continue
        seen.add(item)
        claim, sentiment, support = (_cell(row[col[n]]).lower() for n in ("claim", "sentiment", "support"))
        if not claim:
            problems.append(f"{where}: not labelled (claim is blank)")
            continue
        if claim not in CLAIM_VALUES:
            problems.append(f"{where}: claim {claim!r} is not one of {'/'.join(CLAIM_VALUES)}")
            continue
        if claim == "no":
            labels[item] = {"claim": "no", "sentiment": None, "support": None}
            continue
        issues = []
        for name, value, allowed in (("sentiment", sentiment, SENTIMENT_VALUES),
                                     ("support", support, SUPPORT_VALUES)):
            if not value:
                issues.append(f"missing {name}")
            elif value not in allowed:
                issues.append(f"{name} {value!r} is not one of {'/'.join(allowed)}")
        if issues:
            problems.append(f"{where}: {'; '.join(issues)}")
            continue
        labels[item] = {"claim": "yes", "sentiment": sentiment, "support": support}

    if expected_ids is not None:
        problems += [f"{item}: missing from the sheet" for item in expected_ids if item not in seen]
    if problems:
        raise HumanCheckSheetError(problems)
    return labels


def support_labels(entail: np.ndarray, contra: np.ndarray, threshold: float = 0.5) -> list[str]:
    """The scorer's view of each proposition (columns), as in `matrix_metrics`."""
    max_e, max_c = entail.max(axis=0), contra.max(axis=0)
    return ["supported" if e >= threshold else "contradicted" if c >= threshold else "unsupported"
            for e, c in zip(max_e, max_c)]


def scorer_labels(summaries: Sequence[SampledSummary], collections: Mapping[str, Collection],
                  nli, sentiment, threshold: float = 0.5) -> dict[str, dict]:
    from fairsum.kwsum.scoring import deattributed_matrices

    out: dict[str, dict] = {}
    for s in summaries:
        texts = list(collections[s.collection_id].texts)
        sentences = list(s.sentences)
        entail, contra, entail_d, contra_d, _ = deattributed_matrices(nli, texts, sentences)
        support = support_labels(entail, contra, threshold)
        support_d = support_labels(entail_d, contra_d, threshold)
        if hasattr(sentiment, "labels_both"):
            hybrid, polarity = sentiment.labels_both(sentences)
        else:
            hybrid, polarity = sentiment.labels(sentences), [None] * len(sentences)
        for k, (sup, sup_d, h, p) in enumerate(zip(support, support_d, hybrid, polarity), start=1):
            out[f"{s.summary_id}-{k}"] = {"support": sup, "support_deattr": sup_d, "sentiment": h,
                                          "sentiment_polarity_only": p}
    return out


def binary_support(labels: Mapping[str, str | None]) -> dict[str, str | None]:
    """Supported vs not: the distinction the unsupported rate is built on."""
    return {i: None if v is None else "supported" if v == "supported" else "not_supported"
            for i, v in labels.items()}


def sentiment_view(human: Mapping[str, Mapping]) -> dict[str, str]:
    """Human sentiment on real claims; mixed is left out (the scorer has no such class)."""
    return {i: l["sentiment"] for i, l in human.items()
            if l["claim"] == "yes" and l["sentiment"] != "mixed"}


def support_view(human: Mapping[str, Mapping]) -> dict[str, str]:
    return {i: l["support"] for i, l in human.items() if l["claim"] == "yes"}


@dataclass
class Agreement:
    n: int
    accuracy: float
    alpha: float
    confusion: Counter = field(default_factory=Counter)


def compare(a: Mapping[str, str | None], b: Mapping[str, str | None]) -> Agreement:
    """Accuracy, Krippendorff's alpha (nominal) and (a, b) confusion over shared labelled items."""
    ids = sorted(i for i in set(a) & set(b) if a[i] is not None and b[i] is not None)
    if not ids:
        return Agreement(0, math.nan, math.nan)
    return Agreement(
        n=len(ids),
        accuracy=sum(a[i] == b[i] for i in ids) / len(ids),
        alpha=krippendorff_alpha({i: [a[i], b[i]] for i in ids}),
        confusion=Counter((a[i], b[i]) for i in ids),
    )


def unsupported_rate_gap(human_support: Mapping[str, str], scorer_support: Mapping[str, str],
                         n_boot: int = 2000, seed: int = 0) -> dict:
    """Scorer minus human unsupported rate, with a bootstrap CI that resamples whole summaries.

    Contradicted counts as unsupported on both sides, as in the scorer's
    unsupported rate. A positive difference means the scorer is stricter.
    """
    ids = sorted(set(human_support) & set(scorer_support))
    by_summary: dict[str, list[str]] = defaultdict(list)
    for i in ids:
        by_summary[i.rsplit("-", 1)[0]].append(i)
    keys = sorted(by_summary)

    def rates(groups) -> tuple[float, float]:
        items = [i for g in groups for i in by_summary[g]]
        return (float(np.mean([human_support[i] != "supported" for i in items])),
                float(np.mean([scorer_support[i] != "supported" for i in items])))

    human_rate, scorer_rate = rates(keys)
    rng = np.random.default_rng(seed)
    diffs = []
    for _ in range(n_boot):
        h, s = rates(rng.choice(keys, size=len(keys), replace=True))
        diffs.append(s - h)
    lo, hi = np.percentile(diffs, [2.5, 97.5])
    return {"n_items": len(ids), "n_summaries": len(keys), "human_rate": human_rate,
            "scorer_rate": scorer_rate, "diff": scorer_rate - human_rate,
            "ci_lo": float(lo), "ci_hi": float(hi)}
