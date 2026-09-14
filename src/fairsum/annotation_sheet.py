"""The annotation spreadsheet: export a batch to fill in, read a filled one back.

Kept apart from `annotation` so the schema and its parsing rules need only the
standard library; this module needs openpyxl (`pip install -e ".[annotate]"`).
What a filled sheet may contain is decided by `annotation.parse_sheet_rows` --
this module only moves cells.
"""

from __future__ import annotations

from enum import Enum
from pathlib import Path

import pandas as pd

from fairsum.annotation import (
    N_EVENT_SLOTS,
    Attribution,
    Certainty,
    ReviewAnnotation,
    parse_sheet_rows,
)

# What the annotator reads, what they mark, and the slot fields they fill.
CONTEXT_COLUMNS = ("review_id", "product_id", "condition_name", "text")
STATUS_COLUMNS = ("done", "notes")
SLOT_FIELDS = ("span", "event", "certainty", "attribution")

_WIDTHS = {
    "review_id": 10, "product_id": 12, "condition_name": 12, "text": 70,
    "done": 7, "notes": 24, "span": 28, "event": 18, "certainty": 11, "attribution": 11,
}


def sheet_columns(n_slots: int = N_EVENT_SLOTS) -> list[str]:
    slots = [f"{name}_{i}" for i in range(1, n_slots + 1) for name in SLOT_FIELDS]
    return [*CONTEXT_COLUMNS, *STATUS_COLUMNS, *slots]


def export_sheet(
    batch: pd.DataFrame,
    path: str | Path,
    n_slots: int = N_EVENT_SLOTS,
) -> None:
    """Write `batch` as a sheet with one row per review and empty event slots.

    Refuses to overwrite an existing file: re-running an export over a
    half-annotated sheet would silently destroy the labels.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    path = Path(path)
    if path.exists():
        raise FileExistsError(f"{path} already exists; move it aside to export again")
    missing = [c for c in CONTEXT_COLUMNS if c not in batch.columns]
    if missing:
        raise ValueError(f"batch is missing columns {missing}")

    columns = sheet_columns(n_slots)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "annotations"
    sheet.append(columns)
    for record in batch[list(CONTEXT_COLUMNS)].itertuples(index=False):
        sheet.append([None if pd.isna(value) else value for value in record])
    last_row = sheet.max_row

    dropdowns = {
        "done": DataValidation(type="list", formula1='"yes"', allow_blank=True),
        "certainty": DataValidation(type="list", formula1=_list_formula(Certainty), allow_blank=True),
        "attribution": DataValidation(type="list", formula1=_list_formula(Attribution), allow_blank=True),
    }
    for validation in dropdowns.values():
        sheet.add_data_validation(validation)

    for index, name in enumerate(columns, start=1):
        letter = get_column_letter(index)
        field = name if name in _WIDTHS else name.rsplit("_", 1)[0]
        sheet.column_dimensions[letter].width = _WIDTHS.get(field, 14)
        if field in dropdowns and last_row >= 2:
            dropdowns[field].add(f"{letter}2:{letter}{last_row}")

    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    # Keep the header and the review text in view while scrolling to the slots.
    sheet.freeze_panes = "E2"

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def read_sheet(path: str | Path) -> list[dict[str, object]]:
    """Every non-empty data row as a {header: value} dict; empty cells are None."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook.active.iter_rows(values_only=True)
        header = ["" if h is None else str(h) for h in next(rows, ())]
        return [
            dict(zip(header, values))
            for values in rows
            if any(value is not None for value in values)
        ]
    finally:
        workbook.close()


def load_sheet_annotations(
    path: str | Path,
    n_slots: int = N_EVENT_SLOTS,
) -> list[ReviewAnnotation]:
    """Read a filled sheet; raises `AnnotationSheetError` listing every problem."""
    return parse_sheet_rows(read_sheet(path), n_slots=n_slots)


def _list_formula(options: type[Enum]) -> str:
    return '"' + ",".join(member.value for member in options) + '"'
