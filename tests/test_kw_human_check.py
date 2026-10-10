"""Human check of the scorers (spec §5, human check protocol).

Sampling, blinding, sheet parsing and the agreement numbers are checked with
hand-built generations and labels; model calls are replaced by fixed matrices.
"""

import json
from collections import Counter

import numpy as np
import openpyxl
import pytest

from fairsum.kwsum.collections import Collection
from fairsum.kwsum.human_check import (
    HumanCheckSheetError,
    SHEET_COLUMNS,
    SampledSummary,
    binary_support,
    compare,
    item_ids,
    read_key,
    read_sheet,
    sample_summaries,
    scorer_labels,
    sentiment_view,
    support_labels,
    support_view,
    unsupported_rate_gap,
    web_items,
    web_labels,
    write_key,
    write_reading_pack,
    write_sheet,
)

CONDITIONS = ["direct", "refer", "cot_refer", "refer_oracle", "kw_global", "kw_strat",
              "kw_strat_neg", "key_refer", "key_oracle"]
SUMMARY = ("The drug helped many users sleep through the night. It caused mild headaches "
           "for a few people. The price was considered reasonable by most buyers. Several "
           "reported fast relief of joint pain. Overall the users were satisfied with it.")


def gen(model, dataset, condition, cid, raw=SUMMARY, completion_tokens=80, level="pos75"):
    return {"collection_id": cid, "dataset": dataset, "split": "test", "level": level,
            "condition": condition, "model": model, "raw": raw,
            "completion_tokens": completion_tokens, "status": "ok"}


def grid(models=("m1", "m2"), datasets=("amazon", "drugs"), n_coll=3):
    return [gen(m, d, c, f"{d}-{i}") for m in models for d in datasets
            for c in CONDITIONS for i in range(n_coll)]


def coll(cid, dataset="drugs", polarities=("positive",) * 6 + ("negative",) * 2):
    return Collection(cid, dataset, "test", "pos75", "E", "Sleepwell",
                      tuple(f"r{i}" for i in range(8)),
                      tuple(f"Review text number {i} about sleep." for i in range(8)),
                      tuple(polarities))


class TestSampling:
    def test_even_over_model_and_dataset_and_deterministic(self):
        a = sample_summaries(grid(), n_summaries=8, per_summary=4, seed=42)
        b = sample_summaries(grid(), n_summaries=8, per_summary=4, seed=42)
        assert a == b
        assert Counter((s.model, s.dataset) for s in a) == {
            ("m1", "amazon"): 2, ("m1", "drugs"): 2, ("m2", "amazon"): 2, ("m2", "drugs"): 2}
        assert all(len(s.sentences) == 4 for s in a)

    def test_conditions_are_spread_before_any_repeats(self):
        picked = sample_summaries(grid(models=("m1",), datasets=("drugs",)), n_summaries=9,
                                  per_summary=4, seed=1)
        assert sorted(s.condition for s in picked) == sorted(CONDITIONS)

    def test_distinct_collections_are_preferred(self):
        picked = sample_summaries(grid(models=("m1",), datasets=("drugs",)), n_summaries=3,
                                  per_summary=4, seed=3)
        assert len({s.collection_id for s in picked}) == 3

    def test_truncated_and_short_summaries_are_never_sampled(self):
        gens = [gen("m1", "drugs", "direct", "a", completion_tokens=256),  # hit the cap
                gen("m1", "drugs", "refer", "b", raw="Good product overall. Works well."),
                gen("m1", "drugs", "kw_global", "c")]
        picked = sample_summaries(gens, n_summaries=1, per_summary=4, seed=0, max_new_tokens=256)
        assert [s.collection_id for s in picked] == ["c"]

    def test_too_few_eligible_summaries_is_an_error(self):
        with pytest.raises(ValueError, match="eligible"):
            sample_summaries([gen("m1", "drugs", "direct", "a")], n_summaries=2, per_summary=4)

    def test_sentences_keep_their_order_in_the_summary(self):
        s = sample_summaries(grid(), n_summaries=1, per_summary=4, seed=5)[0]
        positions = [SUMMARY.index(sent) for sent in s.sentences]
        assert positions == sorted(positions)

    def test_summary_ids_are_numbered_and_item_ids_follow(self):
        picked = sample_summaries(grid(), n_summaries=8, per_summary=4, seed=42)
        assert sorted(s.summary_id for s in picked) == [f"S{i:02d}" for i in range(1, 9)]
        assert item_ids(picked)[:4] == ["S01-1", "S01-2", "S01-3", "S01-4"]
        assert len(item_ids(picked)) == 32


class TestFilesAreBlind:
    def test_sheet_and_pack_never_show_model_condition_or_level(self, tmp_path):
        picked = sample_summaries(grid(models=("secretmodel",)), n_summaries=4, per_summary=4)
        write_sheet(picked, tmp_path / "labels.xlsx")
        hidden_polarity = ("POLSENTINEL",) * 8  # review ratings must not leak either
        write_reading_pack(picked, {s.collection_id: coll(s.collection_id, polarities=hidden_polarity)
                                    for s in picked}, tmp_path / "pack.html")
        ws = openpyxl.load_workbook(tmp_path / "labels.xlsx")["labels"]
        cells = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
        html = (tmp_path / "pack.html").read_text(encoding="utf-8")
        for secret in ("secretmodel", "pos75", "kw_strat", "refer_oracle", "POLSENTINEL"):
            assert secret not in cells
            assert secret not in html
        assert [c.value for c in ws[1]] == list(SHEET_COLUMNS)

    def test_pack_shows_all_reviews_and_the_numbered_sentences(self, tmp_path):
        picked = sample_summaries(grid(), n_summaries=1, per_summary=4, seed=5)
        write_reading_pack(picked, {s.collection_id: coll(s.collection_id) for s in picked},
                           tmp_path / "pack.html")
        html = (tmp_path / "pack.html").read_text(encoding="utf-8")
        for i in range(8):
            assert f"Review text number {i} about sleep." in html
        for sent in picked[0].sentences:
            assert sent in html
        assert "S01-4" in html

    def test_key_round_trips(self, tmp_path):
        picked = sample_summaries(grid(), n_summaries=4, per_summary=4)
        write_key(picked, tmp_path / "key.json", seed=42)
        assert read_key(tmp_path / "key.json") == picked


def fill(path, rows):
    """Write labels into a blank sheet: rows maps item_id -> (claim, sentiment, support)."""
    wb = openpyxl.load_workbook(path)
    ws = wb["labels"]
    col = {c.value: i + 1 for i, c in enumerate(ws[1])}
    for r in range(2, ws.max_row + 1):
        item = ws.cell(r, col["item_id"]).value
        if item in rows:
            for name, value in zip(("claim", "sentiment", "support"), rows[item]):
                ws.cell(r, col[name]).value = value
    wb.save(path)


class TestReadSheet:
    @pytest.fixture
    def sheet(self, tmp_path):
        picked = sample_summaries(grid(), n_summaries=1, per_summary=4, seed=5)
        path = tmp_path / "labels.xlsx"
        write_sheet(picked, path)
        return path

    def test_round_trip_with_case_and_spaces_ignored(self, sheet):
        fill(sheet, {"S01-1": ("Yes", " positive ", "SUPPORTED"),
                     "S01-2": ("yes", "negative", "contradicted"),
                     "S01-3": ("yes", "mixed", "unsupported"),
                     "S01-4": ("no", None, None)})
        labels = read_sheet(sheet)
        assert labels["S01-1"] == {"claim": "yes", "sentiment": "positive", "support": "supported"}
        assert labels["S01-3"]["sentiment"] == "mixed"
        assert labels["S01-4"] == {"claim": "no", "sentiment": None, "support": None}

    def test_not_a_claim_ignores_anything_else_on_the_row(self, sheet):
        fill(sheet, {i: ("yes", "neutral", "supported") for i in ("S01-1", "S01-2", "S01-3")}
             | {"S01-4": ("no", "positive", "supported")})
        assert read_sheet(sheet)["S01-4"] == {"claim": "no", "sentiment": None, "support": None}

    def test_every_problem_is_reported_at_once(self, sheet):
        fill(sheet, {"S01-1": ("yes", "happy", "supported"),
                     "S01-2": ("yes", "positive", None),
                     "S01-3": ("maybe", None, None)})  # S01-4 left blank
        with pytest.raises(HumanCheckSheetError) as err:
            read_sheet(sheet)
        text = str(err.value)
        assert len(err.value.problems) == 4
        for fragment in ("S01-1", "'happy'", "S01-2", "support", "S01-3", "'maybe'", "S01-4"):
            assert fragment in text

    def test_missing_items_are_a_problem(self, sheet):
        fill(sheet, {i: ("yes", "neutral", "supported") for i in ("S01-1", "S01-2", "S01-3", "S01-4")})
        with pytest.raises(HumanCheckSheetError, match="S02-1"):
            read_sheet(sheet, expected_ids=["S01-1", "S01-2", "S01-3", "S01-4", "S02-1"])


class TestScorerLabels:
    def test_support_follows_the_scorer_rule(self):
        entail = np.array([[0.9, 0.1, 0.2], [0.1, 0.1, 0.3]])  # reviews x propositions
        contra = np.array([[0.8, 0.7, 0.1], [0.0, 0.2, 0.4]])
        # supported wins over a contradicting review; then contradicted; else unsupported
        assert support_labels(entail, contra, 0.5) == ["supported", "contradicted", "unsupported"]

    def test_scorer_labels_per_item(self):
        picked = sample_summaries(grid(), n_summaries=1, per_summary=4, seed=5)

        class NLI:
            def matrices(self, premises, hypotheses):
                e = np.zeros((len(premises), len(hypotheses)))
                c = np.zeros_like(e)
                e[0, 0], c[0, 1] = 0.9, 0.9
                return e, c

        class Sentiment:
            def labels_both(self, texts):
                return (["positive", "neutral", "negative", "positive"],
                        ["positive", "negative", "negative", "positive"])

        out = scorer_labels(picked, {picked[0].collection_id: coll(picked[0].collection_id)},
                            NLI(), Sentiment())
        assert out["S01-1"] == {"support": "supported", "support_deattr": "supported",
                                "sentiment": "positive", "sentiment_polarity_only": "positive"}
        assert out["S01-2"]["support"] == "contradicted"
        assert out["S01-2"]["sentiment"] == "neutral"
        assert out["S01-2"]["sentiment_polarity_only"] == "negative"
        assert out["S01-3"]["support"] == "unsupported"


    def test_scorer_labels_include_the_attribution_robust_support(self):
        s = SampledSummary("S01", "c1", "m", "direct", "drugs", "pos75", "raw",
                           ("Some reviewers found the scent too strong.", "The price is fair."))

        class FrameSensitiveNLI:  # entails only statements without a reviewer frame
            def matrices(self, premises, hypotheses):
                e = np.array([[0.1 if "reviewers" in h else 0.9 for h in hypotheses] for _ in premises])
                return e, np.zeros_like(e)

        class Sentiment:
            def labels_both(self, texts):
                return ["negative", "neutral"], ["negative", "positive"]

        out = scorer_labels([s], {"c1": coll("c1")}, FrameSensitiveNLI(), Sentiment())
        assert out["S01-1"]["support"] == "unsupported"
        assert out["S01-1"]["support_deattr"] == "supported"   # "I found the scent too strong."
        assert out["S01-2"]["support"] == out["S01-2"]["support_deattr"] == "supported"

    def test_binary_view_collapses_contradicted_into_not_supported(self):
        assert binary_support({"1": "supported", "2": "contradicted", "3": "unsupported", "4": None}) == {
            "1": "supported", "2": "not_supported", "3": "not_supported", "4": None}


class TestAgreement:
    def test_accuracy_alpha_and_confusion(self):
        a = {"1": "positive", "2": "positive", "3": "negative", "4": "neutral", "5": None}
        b = {"1": "positive", "2": "negative", "3": "negative", "4": "neutral", "5": "positive"}
        res = compare(a, b)
        assert res.n == 4  # item 5 has no label on one side
        assert res.accuracy == pytest.approx(0.75)
        assert res.confusion[("positive", "negative")] == 1
        assert res.confusion[("negative", "negative")] == 1
        assert -1 <= res.alpha < 1

    def test_perfect_agreement_gives_alpha_one(self):
        a = {"1": "positive", "2": "negative", "3": "neutral"}
        assert compare(a, dict(a)).alpha == pytest.approx(1.0)

    def test_sentiment_view_drops_non_claims_and_mixed(self):
        human = {"1": {"claim": "yes", "sentiment": "positive", "support": "supported"},
                 "2": {"claim": "yes", "sentiment": "mixed", "support": "supported"},
                 "3": {"claim": "no", "sentiment": None, "support": None}}
        assert sentiment_view(human) == {"1": "positive"}
        assert support_view(human) == {"1": "supported", "2": "supported"}

    def test_unsupported_rate_gap_resamples_whole_summaries(self):
        human = {"S01-1": "supported", "S01-2": "supported", "S02-1": "unsupported",
                 "S02-2": "supported"}
        scorer = {"S01-1": "supported", "S01-2": "unsupported", "S02-1": "unsupported",
                  "S02-2": "contradicted"}
        gap = unsupported_rate_gap(human, scorer, n_boot=500, seed=0)
        assert gap["human_rate"] == pytest.approx(0.25)
        assert gap["scorer_rate"] == pytest.approx(0.75)  # contradicted counts as unsupported
        assert gap["diff"] == pytest.approx(0.5)
        assert gap["ci_lo"] <= 0.5 <= gap["ci_hi"]
        assert gap["n_summaries"] == 2


class TestWebRoute:
    """The labelling page: blinded items go up, labels come back as the same sheets `check` reads."""

    def test_web_items_are_blind_and_complete(self):
        picked = sample_summaries(grid(models=("secretmodel",)), n_summaries=2, per_summary=4)
        items = web_items(picked, {s.collection_id: coll(s.collection_id, polarities=("POLSENTINEL",) * 8)
                                   for s in picked})
        text = json.dumps(items)
        for secret in ("secretmodel", "pos75", "kw_strat", "refer_oracle", "POLSENTINEL",
                       picked[0].collection_id):
            assert secret not in text
        first = items[0]
        assert (first["id"], first["order"], first["topic"]) == (picked[0].summary_id, 1, "Sleepwell")
        assert first["reviews"] == [f"Review text number {i} about sleep." for i in range(8)]
        assert [s["item_id"] for s in first["sentences"]] == [f"{picked[0].summary_id}-{k}" for k in range(1, 5)]
        assert [s["text"] for s in first["sentences"]] == list(picked[0].sentences)

    def test_web_labels_become_a_sheet_that_check_reads(self, tmp_path):
        picked = sample_summaries(grid(), n_summaries=1, per_summary=4, seed=5)
        docs = [{"items": {"S01-1": {"claim": "yes", "sentiment": "positive", "support": "supported",
                                     "notes": "R3 says so"},
                           "S01-2": {"claim": "no", "sentiment": None, "support": None, "notes": ""}}},
                {"items": {"S01-3": {"claim": " Yes", "sentiment": "MIXED", "support": "unsupported"},
                           "S01-4": {"claim": "yes", "sentiment": "negative", "support": "contradicted"},
                           "P1-1": {"claim": "yes", "sentiment": "positive", "support": "supported"}}},
                "not a document", {"items": ["not", "a", "mapping"]}]
        path = tmp_path / "labels_x.xlsx"
        write_sheet(picked, path, labels=web_labels(docs))
        got = read_sheet(path, expected_ids=item_ids(picked))
        assert got["S01-1"] == {"claim": "yes", "sentiment": "positive", "support": "supported"}
        assert got["S01-2"] == {"claim": "no", "sentiment": None, "support": None}
        assert got["S01-3"]["sentiment"] == "mixed"
        assert set(got) == set(item_ids(picked))  # the practice item P1-1 never reaches the sheet
        notes = [c.value for c in openpyxl.load_workbook(path)["labels"]["F"]]
        assert "R3 says so" in notes

    def test_unfinished_web_labels_are_reported_not_defaulted(self, tmp_path):
        picked = sample_summaries(grid(), n_summaries=1, per_summary=4, seed=5)
        labels = web_labels([{"items": {"S01-1": {"claim": "yes", "sentiment": "positive", "support": ""}}}])
        path = tmp_path / "labels_x.xlsx"
        write_sheet(picked, path, labels=labels)
        with pytest.raises(HumanCheckSheetError) as err:
            read_sheet(path, expected_ids=item_ids(picked))
        assert len(err.value.problems) == 4  # S01-1 lacks support; S01-2..4 not labelled
